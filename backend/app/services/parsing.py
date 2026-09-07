import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader


DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


class UnsupportedDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    media_type: str


def normalize_text(text: str) -> str:
    normalized_newlines = text.replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = re.split(r"\n\s*\n", normalized_newlines)
    cleaned = [re.sub(r"\s+", " ", paragraph).strip() for paragraph in paragraphs]
    return "\n\n".join(paragraph for paragraph in cleaned if paragraph)


def parse_docx(path: Path) -> str:
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" ".join(cells))
    return "\n\n".join(blocks)


def parse_pdf(path: Path) -> str:
    reader = PdfReader(path)
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def parse_document(path: Path, media_type: str) -> ParsedDocument:
    if media_type in {"text/plain", "text/markdown"}:
        text = path.read_text(encoding="utf-8-sig")
    elif media_type == DOCX_MEDIA_TYPE:
        text = parse_docx(path)
    elif media_type == "application/pdf":
        text = parse_pdf(path)
    else:
        raise UnsupportedDocumentError(f"Unsupported document type: {media_type}")

    return ParsedDocument(text=normalize_text(text), media_type=media_type)
