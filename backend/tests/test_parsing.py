from pathlib import Path

import pytest

from app.services.parsing import UnsupportedDocumentError, parse_document


@pytest.fixture
def fixture_dir(tmp_path: Path) -> Path:
    source_dir = Path(__file__).parent / "fixtures"
    for name in ("sample.txt", "sample.md"):
        (tmp_path / name).write_bytes((source_dir / name).read_bytes())

    from docx import Document

    document = Document()
    document.add_heading("DocPilot product guide")
    document.add_paragraph("The refund window is 30 days.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Support"
    table.cell(0, 1).text = "support@example.test"
    document.save(tmp_path / "sample.docx")

    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): font_ref}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (The refund window is 30 days.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    with (tmp_path / "sample.pdf").open("wb") as output:
        writer.write(output)

    return tmp_path


@pytest.mark.parametrize(
    ("fixture_name", "media_type"),
    [
        ("sample.txt", "text/plain"),
        ("sample.md", "text/markdown"),
        (
            "sample.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        ("sample.pdf", "application/pdf"),
    ],
)
def test_parse_supported_document(
    fixture_dir: Path, fixture_name: str, media_type: str
):
    parsed = parse_document(fixture_dir / fixture_name, media_type)

    assert "The refund window is 30 days." in parsed.text


def test_parse_docx_includes_table_text(fixture_dir: Path):
    parsed = parse_document(
        fixture_dir / "sample.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert "Support support@example.test" in parsed.text


def test_parse_rejects_executable(tmp_path: Path):
    path = tmp_path / "bad.exe"
    path.write_bytes(b"MZ")

    with pytest.raises(UnsupportedDocumentError):
        parse_document(path, "application/octet-stream")
